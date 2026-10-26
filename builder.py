"""mxene_au_box.py with its variables as inputs. MXenes and metals are the rows of the spreadsheet.

    python builder.py [inputs.json]      writes POSCAR, POSCAR_<mxene>, POSCAR_<metal> and outputs.json
"""
import hashlib
import json
import math
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
from pymatgen.core import Element, Lattice, Structure
from pymatgen.symmetry.analyzer import SpacegroupAnalyzer

SPREADSHEET = Path(__file__).resolve().parent / "data" / "MXene_FCC_L12_strain_matrix.xlsx"

DEFAULTS = dict(
    mxene="V2C",         # MXenes sheet, Formula
    stacking="ABC",      # ABC = AβC (octahedral), ABA = AβA (trigonal prismatic)
    mxene_rotation=0,    # degrees about z: 0, or 180 to turn the MXene over in-plane (B ↔ C, AβC → AγB)
    vacuum=20.0,         # Å, below the MXene only
    metal="Au",          # Metals sheet, Symbol
    metal_layers=4,      # metal (111) atomic planes
    metal_site="auto",   # A, B or C site of the first metal (111) layer; auto = the one free of the top two MXene layers
    metal_stacking="ABC",  # A → B → C upward (continues an MXene at 0°); ACB = A → C → B (continues one at 180°)
    gap=2.3,             # Å, top MXene layer to first metal layer
    shift=[0.0, 0.0],    # Å, in-plane shift of the metal along x and y
    tol=0.02,            # largest |strain| allowed on the MXene along x and along y
    max_replicas=10,     # largest number of metal replicas tried per direction
)

SITES = {"A": (0, 0), "B": (0, 1 / 3), "C": (1 / 2, 1 / 6)}    # hexagonal A, B, C sites in the a × √3a cell
D_MX = 0.385                                                  # M–X layer spacing / a
BOND_MO = 2.0                                                 # Å, M–O bond that sets the O layer height
SPACE_GROUP = {"FCC": ("Fm-3m", 225), "L12": ("Pm-3m", 221)}
ORIENT = [[1, -1, 0], [1, 1, -2], [1, 1, 1]]                  # cubic → x ∥ [1-10], y ∥ [11-2], z ∥ [111]
NEXT = {"ABC": {"A": "B", "B": "C", "C": "A"}, "ACB": {"A": "C", "C": "B", "B": "A"}}


def site_of(fx, fy):
    """A, B or C for a point of a hexagonal net, in fractions of its a × √3a cell."""
    for name, (x, y) in SITES.items():
        for dx, dy in [(0, 0), (0.5, 0.5)]:
            ex, ey = (fx - x - dx) % 1, (fy - y - dy) % 1
            if min(ex, 1 - ex) < 1e-3 and min(ey, 1 - ey) < 1e-3:
                return name
    raise ValueError(f"({fx:.4f}, {fy:.4f}) is not on an A, B or C site")


def mxene_layers(formula):
    """Bottom → top elements of an ordered MXene (Ti3C2, Mo2TiC2, Ti2VCNO2); None for solid solutions and Tx."""
    if not re.fullmatch(r"([A-Z][a-z]?\d*)+", formula):
        return None
    tokens = [(e, int(n or 1)) for e, n in re.findall(r"([A-Z][a-z]?)(\d*)", formula)]
    if not all(Element.is_valid_symbol(e) for e, _ in tokens):
        return None
    metals = [(e, n) for e, n in tokens if e not in ("C", "N", "O")]
    xs = [e for e, n in tokens if e in ("C", "N") for _ in range(n)]
    outer = next((e for e, n in metals if n == 2 and len(metals) > 1), metals[0][0])    # outer metal: the one with count 2
    inner = [e for e, n in metals for _ in range(n) if e != outer] if len(metals) > 1 else [outer] * (len(xs) - 1)
    ms = [outer] + inner + [outer]
    core = [ms[0]] + [e for pair in zip(xs, ms[1:]) for e in pair]
    o = ["O"] if dict(tokens).get("O") else []
    return o + core + o


def read_spreadsheet(path=SPREADSHEET):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mxenes, metals = {}, {}
    for f, family, termination, a, source, *_ in book["MXenes"].iter_rows(min_row=3, values_only=True):
        if f and isinstance(a, (int, float)) and mxene_layers(str(f)):
            mxenes[f] = dict(formula=f, family=family, termination=termination, a=float(a), source=source)
    for symbol, name, structure, a, a111, source, *_ in book["Metals"].iter_rows(min_row=3, values_only=True):
        if symbol and isinstance(a, (int, float)) and re.fullmatch(r"[A-Z][a-z]?(3[A-Z][a-z]?)?", symbol):
            if not isinstance(a111, (int, float)) or abs(a111 - a / math.sqrt(2)) > 1e-3:
                raise ValueError(f"{symbol}: a_111 = {a111} is not a_cubic/√2 = {a / math.sqrt(2):.4f}")
            metals[symbol] = dict(symbol=symbol, name=name, structure=structure, a=float(a), a111=float(a111),
                                  source=source)
    book.close()
    return mxenes, metals


MXENES, METALS = read_spreadsheet()
SHA256 = hashlib.sha256(SPREADSHEET.read_bytes()).hexdigest()


def build(inputs):
    """inputs: dict, missing keys take DEFAULTS. Returns (interface Structure, outputs dict)."""
    p = {**DEFAULTS, **inputs}
    if p["mxene"] not in MXENES or p["metal"] not in METALS:
        raise ValueError(f"{p['mxene']} or {p['metal']} is not a row of {SPREADSHEET.name}")
    if p["metal_site"] not in ("auto", *SITES) or p["metal_stacking"] not in NEXT:
        raise ValueError("metal_site is auto, A, B or C; metal_stacking is ABC or ACB")
    if p["mxene_rotation"] not in (0, 180):
        raise ValueError("mxene_rotation is 0 or 180")

    # ---------------- MXene: conventional (orthogonal) cell, vacuum only below ----------------
    row = MXENES[p["mxene"]]
    a, vac = row["a"], p["vacuum"]
    layers = mxene_layers(row["formula"])

    # AβC: sites cycle A → B → C; AβA: metals on one site, C/N on the next. O sits on the site free of its two
    # nearest layers. Labels are then shifted (a rigid in-plane translation) so the bottom layer is on A.
    # Turning the MXene 180° about z keeps A and swaps B ↔ C: (0, 1/3) → (0, -1/3) ≡ (1/2, 1/6).
    core = [i for i, e in enumerate(layers) if e != "O"]
    centre = (core[0] + core[-1]) // 2
    sites = [None if e == "O" else "ABC"[(i - centre) % 3] if p["stacking"] == "ABC" else "B" if e in ("C", "N") else "A"
             for i, e in enumerate(layers)]
    if layers[0] == "O":
        sites[0] = ({"A", "B", "C"} - {sites[1], sites[2]}).pop()
    if layers[-1] == "O":
        sites[-1] = ({"A", "B", "C"} - {sites[-2], sites[-3]}).pop()
    k = "ABC".index(sites[0])
    sites = ["ABC"[("ABC".index(s) - k) % 3] for s in sites]
    if p["mxene_rotation"] == 180:
        sites = [{"A": "A", "B": "C", "C": "B"}[s] for s in sites]

    heights = [0.0]                                    # above the bottom layer, Å
    for lower, upper in zip(layers, layers[1:]):
        heights.append(heights[-1] + (math.sqrt(BOND_MO**2 - a**2 / 3) if "O" in (lower, upper) else D_MX * a))

    c = heights[-1] + vac                              # top layer sits exactly at z = c
    lattice = Lattice.orthorhombic(a, np.sqrt(3) * a, c)
    species, coords, layer = [], [], []
    for i, (s, site, h) in enumerate(zip(layers, sites, heights)):
        x, y = SITES[site]
        for dx, dy in [(0, 0), (0.5, 0.5)]:            # the rectangular cell holds two copies of each layer
            species.append(s)
            coords.append([(x + dx) % 1, (y + dy) % 1, (vac + h) / c])
            layer.append(i)
    mxene = Structure(lattice, species, coords, to_unit_cell=False, site_properties={"layer": layer})

    # ---------------- Metal: conventional cubic cell from its space group ----------------
    metal_row = METALS[p["metal"]]
    a_m = metal_row["a"]
    symbol, number = SPACE_GROUP[metal_row["structure"]]
    if metal_row["structure"] == "FCC":                # metal at 4a, 4 atoms
        metal = Structure.from_spacegroup(symbol, Lattice.cubic(a_m), [metal_row["symbol"]], [[0, 0, 0]])
    else:                                              # L12 A3B: B at 1a, A at 3c
        A, B = re.fullmatch(r"([A-Z][a-z]?)3([A-Z][a-z]?)", metal_row["symbol"]).groups()
        metal = Structure.from_spacegroup(symbol, Lattice.cubic(a_m), [B, A], [[0, 0, 0], [0, 0.5, 0.5]])
    assert SpacegroupAnalyzer(metal, symprec=1e-3).get_space_group_number() == number

    # ---------------- Metal (111) as stacked hexagonal layers, in the same a × √3a form as the MXene ----------------
    # The [1-10], [11-2], [111] supercell is orthogonal and right-handed, so an orthorhombic lattice with the same
    # fractional coordinates is the cubic crystal rotated to x ∥ [1-10], y ∥ [11-2], z ∥ [111]: 3 layers, 8 atoms each.
    a111, d111 = metal_row["a111"], a_m / math.sqrt(3)
    rotated = metal.copy()
    rotated.make_supercell(ORIENT)
    frac = np.round(rotated.frac_coords, 8) % 1           # rounded: -1e-17 must wrap to 0, not to 1
    if metal_row["structure"] == "FCC":                # one atom of the hexagonal net per 1/2 × 1/2: the a111 × √3a111 cell
        keep = (frac[:, 0] < 0.5 - 1e-6) & (frac[:, 1] < 0.5 - 1e-6)
        net = 1                                        # cell = net × (a111 × √3a111)
        species, frac = [rotated[k].species_string for k in np.flatnonzero(keep)], frac[keep] * [2, 2, 1]
    else:                                              # L12 keeps 2a111 × 2√3a111: B forms a 2 × 2 superlattice
        net = 2
        species = [s.species_string for s in rotated]
    assert np.allclose(rotated.lattice.abc, [2 * a111, 2 * np.sqrt(3) * a111, 3 * d111], atol=1e-3)
    oriented = Structure(Lattice.orthorhombic(net * a111, net * np.sqrt(3) * a111, 3 * d111), species, frac)
    assert len(oriented) == 6 * net**2

    # Bottom layer onto site A, then mirror y (B ↔ C) if the layers climb in the other sense.
    nxt = NEXT[p["metal_stacking"]]
    origin = next(f for f in oriented.frac_coords if f[2] < 1e-6)
    frac = np.round(oriented.frac_coords - [origin[0], origin[1], 0], 8) % 1
    second = next(f for f in frac if abs(f[2] - 1 / 3) < 1e-6)
    mirror = site_of(second[0] * net, second[1] * net) != nxt["A"]
    if mirror:
        frac[:, 1] = -frac[:, 1] % 1
    oriented = Structure(oriented.lattice, oriented.species, frac)
    for f in oriented.frac_coords:                     # layers 0, 1, 2 on A, then the stacking
        assert site_of(f[0] * net, f[1] * net) == ["A", nxt["A"], nxt[nxt["A"]]][round(f[2] * 3) % 3]

    # ---------------- In-plane replicas giving (approximately) the same box ----------------
    # The metal is the substrate: the box is a whole number of metal cells and the MXene takes the strain.
    replicas = []                                      # per direction: (MXene replicas, metal replicas, strain on MXene)
    for L_mx, L_m in zip(mxene.lattice.abc[:2], oriented.lattice.abc[:2]):
        for n in range(1, p["max_replicas"] + 1):
            m = max(1, round(n * L_m / L_mx))
            strain = n * L_m / (m * L_mx) - 1
            if abs(strain) <= p["tol"]:
                replicas.append((m, n, strain))
                break
        else:
            raise ValueError(f"No match within {p['tol']:.1%} up to {p['max_replicas']} replicas; raise the tolerance or max replicas")
    (mx_x, m_x, eps_x), (mx_y, m_y, eps_y) = replicas

    # ---------------- Location: metal (111) on MXene (0001), metal [1-10] along x ----------------
    gap, shift, n_layers = p["gap"], p["shift"], p["metal_layers"]
    first = ({"A", "B", "C"} - {sites[-1], sites[-2]}).pop() if p["metal_site"] == "auto" else p["metal_site"]
    metal_sites = [first]
    while len(metal_sites) < n_layers:
        metal_sites.append(nxt[metal_sites[-1]])
    metal_super = oriented.copy()
    metal_super.make_supercell([m_x, m_y, math.ceil(n_layers / 3)])       # three (111) planes per oriented cell
    metal_super.remove_sites([k for k, s in enumerate(metal_super) if s.coords[2] > (n_layers - 1) * d111 + 1e-3])

    box_x, box_y = m_x * oriented.lattice.a, m_y * oriented.lattice.b
    a_s, b_s = box_x / mx_x, box_y / mx_y              # strained MXene cell; its layer heights stay unstrained
    z_top = c + gap + metal_super.cart_coords[:, 2].max()                  # vacuum only below the MXene
    interface_lattice = Lattice.orthorhombic(box_x, box_y, z_top)

    species, coords, layer, part = [], [], [], []
    for i in range(mx_x):                              # MXene replicas, copied by hand: make_supercell
        for j in range(mx_y):                          # would wrap the top layer from z = c to z = 0
            for s in mxene:                            # same fractional x, y in the strained cell = strained MXene
                species.append(s.species_string)
                coords.append([(i + s.frac_coords[0]) * a_s, (j + s.frac_coords[1]) * b_s, s.coords[2]])
                layer.append(s.properties["layer"])
                part.append("mxene")

    # Bottom metal layer is on A at the origin; move it to the first site of the MXene grid, then add the shift.
    offset = np.array(SITES[first]) * [a_s, b_s] + shift
    for s in metal_super:                              # the box is whole metal cells: the metal keeps its lattice
        species.append(s.species_string)
        x = (s.frac_coords[0] + offset[0] / box_x) % 1 * box_x
        y = (s.frac_coords[1] + offset[1] / box_y) % 1 * box_y
        coords.append([x, y, c + gap + s.coords[2]])
        layer.append(len(layers) + round(s.coords[2] / d111))
        part.append("metal")

    interface = Structure(interface_lattice, species, coords, coords_are_cartesian=True, to_unit_cell=False,
                          site_properties={"layer": layer, "part": part})

    # ---------------- Outputs ----------------
    z = interface.cart_coords[:, 2]
    stack = []
    for i in range(len(layers) + n_layers):
        on = [k for k, L in enumerate(layer) if L == i]
        names = [species[k] for k in on]
        stack.append(dict(index=i, part=part[on[0]], site=sites[i] if i < len(layers) else metal_sites[i - len(layers)],
                          species={e: names.count(e) for e in dict.fromkeys(names)}, z=round(float(z[on].mean()), 4)))

    r = lambda v, n=4: round(float(v), n)
    turn = -1 if p["mxene_rotation"] == 180 else 1                         # the MXene's a and b turn with it
    # Cubic cell vectors in the interface frame (rows), for the (hkl) of the viewer.
    rotation = np.array(ORIENT, dtype=float) / np.linalg.norm(ORIENT, axis=1)[:, None]
    if mirror:
        rotation[1] *= -1
    outputs = dict(
        inputs=p,
        mxene=dict(**row, stacking=p["stacking"], rotation=p["mxene_rotation"], layers=layers, sites=sites,
                   cell=[r(a), r(np.sqrt(3) * a), r(c)], strained_cell=[r(a_s), r(b_s), r(c)],
                   matrix=(np.diag([turn * a_s, turn * b_s, c]).round(4) + 0.0).tolist(),  # strained cell vectors, interface frame
                   thickness=r(heights[-1])),
        metal=dict(**metal_row, space_group=f"{symbol} ({number})", layers=n_layers, d111=r(d111),
                   orientation="(111) ∥ MXene (0001), [1-10] ∥ x, [11-2] ∥ y", stacking=p["metal_stacking"],
                   sites=metal_sites, cell=[r(v) for v in oriented.lattice.abc],
                   matrix=np.round(a_m * rotation.T + 0.0, 6).tolist()),         # cubic cell vectors, interface frame
        match=dict(x=dict(mxene=mx_x, metal=m_x, length=r(box_x), strain=r(eps_x, 6)),
                   y=dict(mxene=mx_y, metal=m_y, length=r(box_y), strain=r(eps_y, 6))),
        interface=dict(formula=interface.composition.formula, n_atoms=len(interface), box=[r(box_x), r(box_y), r(z_top)],
                       layers=stack),
        spreadsheet=dict(file=SPREADSHEET.name, sha256=SHA256),
        files={"POSCAR": interface.get_sorted_structure().to(fmt="poscar"),
               f"POSCAR_{p['mxene']}": mxene.to(fmt="poscar"),
               f"POSCAR_{p['metal']}": metal.to(fmt="poscar")},
    )
    return interface, outputs


if __name__ == "__main__":
    inputs = json.loads(Path(sys.argv[1]).read_text()) if len(sys.argv) > 1 else {}
    _, outputs = build(inputs)
    for name, text in outputs["files"].items():
        Path(name).write_text(text)
    Path("outputs.json").write_text(json.dumps(outputs, indent=2) + "\n")
    print(outputs["interface"]["formula"], outputs["match"])
